"""Read royalty statement files (CSV/TSV/XLSX) with heterogeneous layouts.

The parser:
  * detects the header row (statements often have preamble/title rows),
  * maps varied column names onto canonical fields via an alias table,
  * parses flexible date/period formats (full dates, "2025-03", "Mar 2025", "Q1 2025"),
  * parses amounts with currency symbols, thousands separators and parentheses,
  * keeps every raw row so each transaction is traceable to its origin.

Files are STREAMED: `StatementReader` yields rows lazily with bounded memory,
so multi-hundred-megabyte statements ingest without loading into RAM. CSV/TSV
go through the stdlib csv module; .xlsx/.xlsm through openpyxl in read-only
mode; legacy .xls (small by nature) through pandas.
"""

from __future__ import annotations

import calendar
import csv
import hashlib
import json
import re
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from itertools import chain
from pathlib import Path
from typing import Iterator

from dateutil import parser as dateutil_parser

from .categorize import UNCATEGORIZED, categorize_row

# ---------------------------------------------------------------------------
# Column alias table
# ---------------------------------------------------------------------------

def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h).lower())


# field -> list of aliases, ordered by preference (earlier = preferred when
# several columns of the same field exist, e.g. "net amount" beats "gross amount").
_FIELD_ALIASES: dict[str, list[str]] = {
    "amount": [
        "netamount", "netroyalty", "netroyalties", "netrevenue", "netearnings",
        "netincome", "netpayable", "netamountpayable", "netpayout", "amountdue",
        "amountpayable", "royaltyamount", "royaltiesearned", "royaltyearned",
        "yourshare", "payableamount", "netdistamount", "netdistribution",
        "netpayment", "paymentamount", "amountpaid", "royaltyvalue",
        "distamount", "distributionamount", "earnings", "royalty",
        "royalties", "payout", "revenue", "income", "amount", "total",
        "totalamount", "totalearnings", "totalroyalty", "totalroyalties",
        "value", "netdue", "grossamount", "grossroyalty",
    ],
    "date": [
        "transactiondate", "statementdate", "paymentdate", "saledate",
        "activitydate", "date", "statementperiod", "accountingperiod",
        "royaltyperiod", "salesperiod", "reportingperiod", "distributionperiod",
        "performanceperiod", "period", "periodend", "periodending",
        "salemonth", "reportingmonth", "activitymonth", "month", "quarter",
        "year",
    ],
    "income_type": [
        "incometype", "royaltytype", "righttype", "rightstype", "revenuetype",
        "incomecategory", "royaltycategory", "revenuecategory", "earningstype",
        "usetype", "usagetype", "type", "category", "incomesource",
        "revenuesource", "royaltysource", "configuration",
    ],
    "source": [
        "source", "payor", "payer", "society", "collectionsociety", "platform",
        "store", "storename", "service", "dsp", "retailer", "distributor",
        "channel", "saletype", "outlet", "partner", "licensee", "territorysource",
    ],
    "track": [
        "tracktitle", "songtitle", "worktitle", "recordingtitle", "releasetitle",
        "title", "track", "song", "work", "recording", "release", "trackname",
        "songname", "assettitle", "project", "isrctitle",
    ],
    "artist": [
        "artistname", "artist", "band", "performer", "actname",
    ],
    "description": [
        "description", "details", "notes", "memo", "usage", "lineitem",
        "transactiondescription", "comment", "comments",
    ],
    "currency": [
        "currency", "currencycode", "curr", "ccy",
    ],
}

# normalized alias -> (field, preference_rank)
_ALIAS_LOOKUP: dict[str, tuple[str, int]] = {}
for _field, _aliases in _FIELD_ALIASES.items():
    for _rank, _alias in enumerate(_aliases):
        _ALIAS_LOOKUP.setdefault(_alias, (_field, _rank))

_CURRENCY_SUFFIXES = (
    "gbp", "usd", "eur", "cad", "aud", "jpy", "chf", "sek", "nok", "dkk",
    "nzd", "brl", "mxn", "sterling", "dollars", "euros", "pounds",
)


def _lookup_alias(normalized: str) -> tuple[str, int] | None:
    """Alias lookup that also tolerates currency-qualified headers, so
    'Earnings (USD)', 'Amount GBP' and 'Net Amount EUR' all resolve."""
    hit = _ALIAS_LOOKUP.get(normalized)
    if hit:
        return hit
    for suffix in _CURRENCY_SUFFIXES:
        if normalized.endswith(suffix) and len(normalized) > len(suffix):
            return _ALIAS_LOOKUP.get(normalized[: -len(suffix)])
    return None


# ---------------------------------------------------------------------------
# Value parsing
# ---------------------------------------------------------------------------

_CURRENCY_SYMBOLS = {"$": "USD", "€": "EUR", "£": "GBP", "¥": "JPY"}
_AMOUNT_CLEAN_RE = re.compile(r"[^\d.,\-()]")


@lru_cache(maxsize=65536)  # statement cells repeat heavily; caching is a big win
def parse_amount(raw: str | float | int | None) -> Decimal | None:
    """Parse '$1,234.56', '(12.50)', '1.234,56', '0.003' etc. into a Decimal."""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s or s.lower() in {"nan", "none", "n/a", "-", "--"}:
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative = True
        s = s[1:-1]
    s = _AMOUNT_CLEAN_RE.sub("", s).strip()
    if s.startswith("-"):
        negative = True
        s = s[1:]
    s = s.replace("(", "").replace(")", "")
    if not s:
        return None
    # Decide decimal separator: if both present, the right-most one is decimal.
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        # "1,234" -> thousands; "12,5" -> decimal comma
        tail = s.split(",")[-1]
        if len(tail) == 3 and len(s.split(",")[0]) > 0:
            s = s.replace(",", "")
        else:
            s = s.replace(",", ".")
    try:
        value = Decimal(s)
    except InvalidOperation:
        return None
    return -value if negative else value


def detect_currency(raw_amount: str | None, header: str | None) -> str | None:
    """Pull a currency hint from the amount cell ('$12') or header ('Earnings (USD)')."""
    for text in (raw_amount, header):
        if not text:
            continue
        text = str(text)
        m = re.search(r"\b(USD|EUR|GBP|CAD|AUD|JPY|SEK|NOK|DKK|CHF|NZD|BRL|MXN)\b",
                      text.upper())
        if m:
            return m.group(1)
        for sym, code in _CURRENCY_SYMBOLS.items():
            if sym in text:
                return code
    return None


_MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
_MONTHS.update({m.lower(): i for i, m in enumerate(calendar.month_abbr) if m})


def _end_of_month(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


@lru_cache(maxsize=16384)  # dates/periods repeat across a statement's rows
def parse_statement_date(raw: str | None) -> date | None:
    """Parse a transaction date or statement period into a date.

    Periods (months, quarters, years) resolve to the period's last day, so a
    'March 2025' line lands in calendar year 2025 and in any LTM window that
    includes end of March 2025.
    """
    if raw is None:
        return None
    s = str(raw).strip()
    if not s or s.lower() in {"nan", "none", "n/a"}:
        return None

    # Quarter: "Q1 2025", "2025 Q1", "2025-Q1"
    m = re.fullmatch(r"(?i)q([1-4])[\s\-/]*(\d{4})", s) or \
        re.fullmatch(r"(?i)(\d{4})[\s\-/]*q([1-4])", s)
    if m:
        a, b = m.group(1), m.group(2)
        year, q = (int(b), int(a)) if len(a) == 1 else (int(a), int(b))
        return _end_of_month(year, q * 3)

    # Year only: "2025"
    if re.fullmatch(r"(19|20)\d{2}", s):
        return date(int(s), 12, 31)

    # Numeric year-month: "2025-03", "2025/03", "202503", "03/2025"
    m = re.fullmatch(r"(\d{4})[\-/.](\d{1,2})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return _end_of_month(int(m.group(1)), int(m.group(2)))
    m = re.fullmatch(r"(\d{1,2})[\-/.](\d{4})", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return _end_of_month(int(m.group(2)), int(m.group(1)))
    m = re.fullmatch(r"(\d{4})(\d{2})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return _end_of_month(int(m.group(1)), int(m.group(2)))

    # Month-name periods: "Mar 2025", "March 2025", "Mar-25", "2025 March"
    m = re.fullmatch(r"(?i)([a-z]{3,9})[\s\-/,]+(\d{2,4})", s)
    if m and m.group(1).lower() in _MONTHS:
        year = int(m.group(2))
        if year < 100:
            year += 2000
        return _end_of_month(year, _MONTHS[m.group(1).lower()])
    m = re.fullmatch(r"(?i)(\d{4})[\s\-/,]+([a-z]{3,9})", s)
    if m and m.group(2).lower() in _MONTHS:
        return _end_of_month(int(m.group(1)), _MONTHS[m.group(2).lower()])

    # Full dates - dateutil handles ISO, US, and most exports.
    try:
        return dateutil_parser.parse(s, dayfirst=False).date()
    except (ValueError, OverflowError):
        pass
    try:
        return dateutil_parser.parse(s, dayfirst=True).date()
    except (ValueError, OverflowError):
        return None


# ---------------------------------------------------------------------------
# File reading (streaming)
# ---------------------------------------------------------------------------

@dataclass
class ParsedRow:
    row_number: int              # 1-based data row number in the original file
    txn_date: date | None
    amount: Decimal
    currency: str | None
    category: str
    category_reason: str
    income_type: str | None
    source: str | None
    track: str | None
    artist: str | None
    description: str | None
    raw: dict = field(default_factory=dict)

    @property
    def fingerprint(self) -> str:
        """Content hash used for cross-file duplicate detection.

        Built from the raw original row values (not the interpreted fields), so
        two different files carrying the identical statement line collide, while
        rows that differ in any original column do not.
        """
        payload = json.dumps(
            {k: v for k, v in self.raw.items()}, sort_keys=True, ensure_ascii=False
        )
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()


@dataclass
class ParseResult:
    rows: list[ParsedRow]
    skipped: list[tuple[int, str]]   # (row_number, reason) for rows we could not use
    column_map: dict[str, str]       # canonical field -> original column name
    warnings: list[str]


_TOTAL_ROW_RE = re.compile(r"(?i)^\s*(grand\s+)?(sub)?total\b")

MAX_SKIPPED_SAMPLES = 200   # keep at most this many skip reasons in memory


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() == "nan" else s


def _delimited_rows(path: Path) -> Iterator[list[str]]:
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
        sample = fh.read(64 * 1024)
        fh.seek(0)
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
        for row in csv.reader(fh, delimiter=delimiter):
            yield [c.strip() for c in row]


def _xlsx_sheet_rows(ws) -> Iterator[list[str]]:
    for row in ws.iter_rows(values_only=True):
        yield [_cell_to_str(c) for c in row]


def _sheet_iterators(path: Path) -> tuple[list[tuple[str, Iterator[list[str]]]], callable]:
    """Return ((sheet_name, row_iterator), ...) for every table in the file,
    plus a close callback. CSVs have one unnamed sheet; workbooks expose every
    worksheet so data hiding behind a cover sheet is still found."""
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        return [(ws.title, _xlsx_sheet_rows(ws)) for ws in wb.worksheets], wb.close
    if suffix == ".xls":
        # Legacy format with a hard 65k-row cap - small enough to load whole.
        import pandas as pd

        frames = pd.read_excel(path, header=None, dtype=str, sheet_name=None)
        sheets = []
        for name, df in frames.items():
            rows = [
                ["" if pd.isna(c) else _cell_to_str(c) for c in row]
                for _, row in df.iterrows()
            ]
            sheets.append((str(name), iter(rows)))
        return sheets, (lambda: None)
    return [("", _delimited_rows(path))], (lambda: None)


def _preview(rows: list[list[str]], max_rows: int = 4, max_cells: int = 8,
             cell_width: int = 24) -> str:
    """Compact preview of a file's first rows, for unrecognized-format errors."""
    lines = []
    for cells in rows:
        if not any(cells):
            continue
        shown = [
            (c[: cell_width - 1] + "…") if len(c) > cell_width else c
            for c in cells[:max_cells]
        ]
        suffix = f" …+{len(cells) - max_cells} cols" if len(cells) > max_cells else ""
        lines.append(" | ".join(shown) + suffix)
        if len(lines) >= max_rows:
            break
    return "\n".join(lines) if lines else "(file appears empty)"


def _scan_for_header(rows: list[list[str]]) -> tuple[int, dict[int, tuple[str, int]]]:
    """Locate the header row: the first row where cells map to known fields
    (at minimum an amount column). Returns (row_index, {col_index: (field, rank)}).
    """
    best: tuple[int, dict[int, tuple[str, int]]] | None = None
    for i, cells in enumerate(rows):
        mapping: dict[int, tuple[str, int]] = {}
        for col_idx, cell in enumerate(cells):
            if not cell:
                continue
            hit = _lookup_alias(_norm_header(cell))
            if hit:
                mapping[col_idx] = hit
        fields_found = {f for f, _ in mapping.values()}
        if "amount" in fields_found and len(fields_found) >= 2:
            return i, mapping
        if mapping and best is None:
            best = (i, mapping)
    if best:
        return best
    raise ValueError(
        "Could not find a header row with recognizable columns "
        "(need at least an amount column such as 'Net Amount', 'Earnings' or "
        "'Royalty'). The file starts like this:\n" + _preview(rows)
    )


# ---------------------------------------------------------------------------
# Content-based column inference (for files whose column names we don't know)
# ---------------------------------------------------------------------------

def _looks_like_money(s: str) -> bool:
    """A value that parses as an amount AND carries a money hint (decimal
    separator, currency symbol, negative). Bare integers are rejected so
    quantity/ID/year columns don't masquerade as earnings."""
    if not s or not any(ch.isdigit() for ch in s):
        return False
    if parse_amount(s) is None:
        return False
    return any(ch in s for ch in ".,$€£¥(") or s.startswith("-")


def _plausible_date(s: str) -> bool:
    if not s:
        return False
    # Money-shaped values are never date evidence: dateutil would happily read
    # "10.50" as the time 10:50 today, which poisons column classification.
    if _looks_like_money(s):
        return False
    d = parse_statement_date(s)
    return d is not None and 1990 <= d.year <= 2035


def _column_values(data_rows: list[list[str]], col: int) -> list[str]:
    return [r[col] for r in data_rows if col < len(r) and r[col]]


def _infer_amount_col(data_rows: list[list[str]], exclude: set[int]) -> int | None:
    """Pick the most money-like column: highest fraction of money-shaped
    values, rightmost on ties (statements put the payable amount last)."""
    ncols = max((len(r) for r in data_rows), default=0)
    best: tuple[float, int] | None = None
    for col in range(ncols):
        if col in exclude:
            continue
        vals = _column_values(data_rows, col)
        if len(vals) < 2:
            continue
        money_frac = sum(_looks_like_money(v) for v in vals) / len(vals)
        date_frac = sum(_plausible_date(v) for v in vals) / len(vals)
        if money_frac < 0.7 or date_frac >= 0.7:
            continue
        if best is None or money_frac > best[0] or (money_frac == best[0] and col > best[1]):
            best = (money_frac, col)
    return best[1] if best else None


def _infer_date_col(data_rows: list[list[str]], exclude: set[int]) -> int | None:
    ncols = max((len(r) for r in data_rows), default=0)
    best: tuple[float, int] | None = None
    for col in range(ncols):
        if col in exclude:
            continue
        vals = _column_values(data_rows, col)
        if len(vals) < 2:
            continue
        frac = sum(_plausible_date(v) for v in vals) / len(vals)
        if frac < 0.7:
            continue
        if best is None or frac > best[0]:   # ties keep the leftmost
            best = (frac, col)
    return best[1] if best else None


def _text_columns(data_rows: list[list[str]], exclude: set[int]) -> list[int]:
    """Columns that are mostly words - fed to the categorizer when no named
    type/source/description columns exist."""
    ncols = max((len(r) for r in data_rows), default=0)
    out = []
    for col in range(ncols):
        if col in exclude:
            continue
        vals = _column_values(data_rows, col)
        if len(vals) < 2:
            continue
        alpha = sum(any(ch.isalpha() for ch in v) for v in vals)
        if alpha / len(vals) >= 0.5:
            out.append(col)
    return out


class StatementReader:
    """Streams normalized transaction rows from one statement file.

    Construction reads only the first few rows (header detection); iterating
    parses the rest lazily, so memory stays bounded for arbitrarily large
    files. Bookkeeping (skip counts, currencies seen) accumulates during
    iteration; call `finalize_warnings()` after consuming the iterator.
    """

    HEADER_SCAN_ROWS = 40   # some statements bury the header under long preambles

    def __init__(self, path: str | Path, filename: str | None = None,
                 column_overrides: dict | None = None):
        self.path = Path(path)
        self.filename = filename or self.path.name
        self.skipped_count = 0
        self.skipped_samples: list[tuple[int, str]] = []
        self.currencies: set[str] = set()
        self.warnings: list[str] = []
        self._text_cols: list[int] = []

        sheets, self._close = _sheet_iterators(self.path)
        last_error: ValueError | None = None
        for sheet_name, rows_iter in sheets:
            buffered: list[list[str]] = []
            for row in rows_iter:
                buffered.append(row)
                if len(buffered) >= self.HEADER_SCAN_ROWS:
                    break
            try:
                self.warnings = []
                if column_overrides is not None:
                    self._configure_overrides(buffered, column_overrides)
                else:
                    self._configure(buffered)
            except ValueError as e:
                last_error = e
                continue
            self._raw_iter = rows_iter
            if sheet_name and len(sheets) > 1:
                self.warnings.append(f"Using worksheet '{sheet_name}'.")
            return
        self._close()
        raise last_error or ValueError("The file contains no rows.")

    def _configure_overrides(self, buffered: list[list[str]], overrides: dict) -> None:
        """Apply an externally supplied column mapping (e.g. from Claude)."""
        columns = {f: int(i) for f, i in overrides.get("columns", {}).items()
                   if i is not None}
        if "amount" not in columns:
            raise ValueError("Column mapping has no amount column.")
        header_row = overrides.get("header_row")
        header_idx = header_row if isinstance(header_row, int) and header_row >= 0 else -1
        if header_idx >= len(buffered):
            raise ValueError("Column mapping header row is outside the file.")

        self._headers = buffered[header_idx] if header_idx >= 0 else []
        self._chosen = columns
        self._header_idx = header_idx
        self._pending = buffered[header_idx + 1:]
        amount_idx = columns["amount"]
        self._header_currency = (
            detect_currency(None, self._headers[amount_idx])
            if amount_idx < len(self._headers) else None
        )
        if not any(f in columns for f in ("income_type", "source", "description")):
            self._text_cols = _text_columns(self._pending, exclude=set(columns.values()))
        self.column_map = {f: self._col_name(i) for f, i in columns.items()}
        described = ", ".join(f"{f}: {self._col_name(i)}" for f, i in columns.items())
        self.warnings.append(
            f"Columns identified by Claude ({described}). "
            "Spot-check a few transactions with /trace."
        )
        if "date" not in columns:
            self.warnings.append(
                "No date/period column recognized - rows will need manual review."
            )

    def _col_name(self, idx: int) -> str:
        if idx < len(self._headers) and self._headers[idx]:
            return self._headers[idx]
        return f"column {idx + 1}"

    def _configure(self, buffered: list[list[str]]) -> None:
        """Map columns for one sheet: recognized header names first, content
        inference to fill gaps, full inference when nothing is recognized."""
        try:
            header_idx, col_hits = _scan_for_header(buffered)
        except ValueError:
            self._configure_inferred(buffered)
            return

        headers = buffered[header_idx]
        chosen: dict[str, int] = {}
        chosen_rank: dict[str, int] = {}
        for col_idx, (field_name, rank) in col_hits.items():
            if field_name not in chosen or rank < chosen_rank[field_name]:
                chosen[field_name] = col_idx
                chosen_rank[field_name] = rank

        self._headers = headers
        data_rows = buffered[header_idx + 1:]
        if "amount" not in chosen:
            idx = _infer_amount_col(data_rows, exclude=set(chosen.values()))
            if idx is None:
                raise ValueError(
                    "No amount column recognized in this file. "
                    "The file starts like this:\n" + _preview(buffered)
                )
            chosen["amount"] = idx
            self.warnings.append(
                f"Amount column inferred from the data: '{self._col_name(idx)}'. "
                "Spot-check a few transactions with /trace."
            )
        if "date" not in chosen:
            idx = _infer_date_col(data_rows, exclude=set(chosen.values()))
            if idx is not None:
                chosen["date"] = idx
                self.warnings.append(
                    f"Date column inferred from the data: '{self._col_name(idx)}'."
                )
            else:
                self.warnings.append(
                    "No date/period column recognized - rows will need manual review."
                )
        # When no type/source/description columns exist, categorize from
        # whatever text columns the sheet has.
        if not any(f in chosen for f in ("income_type", "source", "description")):
            self._text_cols = _text_columns(data_rows, exclude=set(chosen.values()))

        self._chosen = chosen
        self._header_idx = header_idx
        self._pending = data_rows
        self._header_currency = detect_currency(None, headers[chosen["amount"]]
                                                if chosen["amount"] < len(headers) else None)
        self.column_map = {f: self._col_name(idx) for f, idx in chosen.items()}

    def _configure_inferred(self, buffered: list[list[str]]) -> None:
        """No recognizable column names anywhere: find where data starts and
        classify columns purely by content (headerless exports like PRS 052)."""
        data_start = next(
            (i for i, row in enumerate(buffered) if any(_looks_like_money(c) for c in row)),
            None,
        )
        if data_start is None:
            raise ValueError(
                "Could not find a header row with recognizable columns "
                "(need at least an amount column such as 'Net Amount', 'Earnings' or "
                "'Royalty'). The file starts like this:\n" + _preview(buffered)
            )
        header_idx = -1
        for i in range(data_start - 1, -1, -1):
            cells = buffered[i]
            if sum(1 for c in cells if c) >= 2 and any(
                any(ch.isalpha() for ch in c) for c in cells
            ):
                header_idx = i
                break

        data_rows = buffered[data_start:]
        amount_idx = _infer_amount_col(data_rows, exclude=set())
        if amount_idx is None:
            raise ValueError(
                "No amount column recognized in this file. "
                "The file starts like this:\n" + _preview(buffered)
            )
        chosen: dict[str, int] = {"amount": amount_idx}
        date_idx = _infer_date_col(data_rows, exclude={amount_idx})
        if date_idx is not None:
            chosen["date"] = date_idx

        self._headers = buffered[header_idx] if header_idx >= 0 else []
        self._chosen = chosen
        self._header_idx = data_start - 1
        self._pending = data_rows
        self._header_currency = None
        self._text_cols = _text_columns(data_rows, exclude=set(chosen.values()))
        self.column_map = {f: self._col_name(idx) for f, idx in chosen.items()}
        described = ", ".join(f"{f}: {self._col_name(i)}" for f, i in chosen.items())
        self.warnings.append(
            f"No recognized column names - columns were inferred from the data "
            f"({described}). Spot-check a few transactions with /trace."
        )
        if "date" not in chosen:
            self.warnings.append(
                "No date/period column recognized - rows will need manual review."
            )

    def _cell(self, cells: list[str], field_name: str) -> str | None:
        idx = self._chosen.get(field_name)
        if idx is None or idx >= len(cells):
            return None
        return cells[idx] or None

    def _skip(self, row_number: int, reason: str) -> None:
        self.skipped_count += 1
        if len(self.skipped_samples) < MAX_SKIPPED_SAMPLES:
            self.skipped_samples.append((row_number, reason))

    def __iter__(self) -> Iterator[ParsedRow]:
        try:
            yield from self._rows()
        finally:
            self._close()

    def _rows(self) -> Iterator[ParsedRow]:
        headers = self._headers
        for offset, cells in enumerate(chain(self._pending, self._raw_iter)):
            row_number = self._header_idx + 2 + offset  # 1-based position in file
            raw: dict[str, str] = {}
            for j in range(max(len(headers), len(cells))):
                key = (headers[j] if j < len(headers) and headers[j] else f"col{j+1}")
                raw[key] = cells[j] if j < len(cells) else ""
            if all(v == "" for v in raw.values()):
                continue

            first_cell = next((v for v in raw.values() if v), "")
            amount = parse_amount(self._cell(cells, "amount"))
            if amount is None:
                self._skip(row_number, "no parseable amount")
                continue
            # Cheap first-letter prefilter: the regex only ever matches values
            # starting with total/subtotal/grand, so skip the rest outright.
            if _TOTAL_ROW_RE.match(first_cell) or any(
                _TOTAL_ROW_RE.match(v)
                for v in raw.values()
                if v and v[0] in "tTgGsS"
            ):
                self._skip(row_number, "looks like a total/subtotal row")
                continue

            txn_date = parse_statement_date(self._cell(cells, "date"))
            income_type = self._cell(cells, "income_type")
            source = self._cell(cells, "source")
            description = self._cell(cells, "description")
            if description is None and self._text_cols:
                # Inferred layout: categorize from the sheet's text columns.
                blob = " ".join(
                    cells[i] for i in self._text_cols if i < len(cells) and cells[i]
                )
                description = blob[:200] or None
            currency = (
                self._cell(cells, "currency")
                or detect_currency(self._cell(cells, "amount"), None)
                or self._header_currency
            )
            if currency:
                currency = currency.upper()
                self.currencies.add(currency)

            category, reason = categorize_row(income_type, source, description, self.filename)
            if txn_date is None:
                # Without a date the transaction cannot be placed in a year or
                # the LTM window, so force manual review regardless of keywords.
                if category != UNCATEGORIZED:
                    reason = f"date missing/unparseable (would be {category}: {reason})"
                else:
                    reason = "date missing/unparseable; " + reason
                category = UNCATEGORIZED

            yield ParsedRow(
                row_number=row_number,
                txn_date=txn_date,
                amount=amount,
                currency=currency,
                category=category,
                category_reason=reason,
                income_type=income_type,
                source=source,
                track=self._cell(cells, "track"),
                artist=self._cell(cells, "artist"),
                description=description,
                raw=raw,
            )

    def finalize_warnings(self) -> list[str]:
        """Warnings including those only known after full iteration."""
        if len(self.currencies) > 1:
            note = (
                "Multiple currencies detected in this file: "
                + ", ".join(sorted(self.currencies))
                + ". Totals do not convert between currencies."
            )
            if note not in self.warnings:
                self.warnings.append(note)
        return self.warnings


def parse_file(path: str | Path, filename: str | None = None) -> ParseResult:
    """Parse a whole statement into memory. Convenience wrapper around
    StatementReader for small files and tests; large-file ingestion streams
    the reader directly instead."""
    reader = StatementReader(path, filename=filename)
    rows = list(reader)
    return ParseResult(
        rows=rows,
        skipped=list(reader.skipped_samples),
        column_map=dict(reader.column_map),
        warnings=reader.finalize_warnings(),
    )


def preview_rows(path: str | Path, limit: int = 25) -> list[list[str]]:
    """First rows of the file's first sheet, for AI layout mapping."""
    sheets, close = _sheet_iterators(Path(path))
    try:
        if not sheets:
            return []
        _, rows_iter = sheets[0]
        out: list[list[str]] = []
        for row in rows_iter:
            out.append(row)
            if len(out) >= limit:
                break
        return out
    finally:
        close()


def layout_signature(rows: list[list[str]], suffix: str) -> str:
    """Structural fingerprint of a statement layout.

    Files exported from the same system share a signature even though their
    values differ, so an AI-derived column mapping can be cached and reused
    across a whole batch of statements."""
    ncols = max((len(r) for r in rows), default=0)
    pattern = []
    for col in range(ncols):
        vals = _column_values(rows, col)
        if not vals:
            pattern.append("e")
        elif sum(_looks_like_money(v) for v in vals) / len(vals) >= 0.5:
            pattern.append("m")
        elif sum(_plausible_date(v) for v in vals) / len(vals) >= 0.5:
            pattern.append("d")
        elif sum(any(ch.isalpha() for ch in v) for v in vals) / len(vals) >= 0.5:
            pattern.append("a")
        else:
            pattern.append("n")
    return f"{suffix.lower()}:{ncols}:{''.join(pattern)}"


def file_sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()
